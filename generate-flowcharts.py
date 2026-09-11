"""ISO-style flowcharts matching the user's Flowchart.docx sample."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = Path(r"C:\Users\Bevin\Documents\NetXScan-flowcharts")
DOCX_PATH = Path(r"C:\Users\Bevin\Documents\NetXScan-Flowcharts.docx")
BG = (0, 0, 0)
FG = (255, 255, 255)
INK = (0, 0, 0)

FONT_PATH = Path(r"C:\Windows\Fonts\arial.ttf")
FONT = ImageFont.truetype(str(FONT_PATH), 13)
FONT_SM = ImageFont.truetype(str(FONT_PATH), 12)
FONT_CAPTION = ImageFont.truetype(str(FONT_PATH), 14)


def wrap(text: str, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur = ""
    for word in words:
        trial = f"{cur} {word}".strip()
        if len(trial) <= width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines or [text]


def draw_text(draw: ImageDraw.ImageDraw, cx: int, cy: int, text: str, width: int, font=FONT) -> None:
    lines = wrap(text, width)
    h = font.size + 2
    y = cy - (len(lines) * h) // 2 + 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        tw = bbox[2] - bbox[0]
        draw.text((cx - tw // 2, y), line, fill=INK, font=font)
        y += h


def oval(draw: ImageDraw.ImageDraw, cx: int, cy: int, text: str, w=128, h=42) -> None:
    x0, y0, x1, y1 = cx - w // 2, cy - h // 2, cx + w // 2, cy + h // 2
    draw.rounded_rectangle((x0, y0, x1, y1), radius=h // 2, fill=FG, outline=FG)
    draw_text(draw, cx, cy, text, 14)


def parallelogram(draw: ImageDraw.ImageDraw, cx: int, cy: int, text: str, w=168, h=70) -> None:
    skew = 22
    pts = [
        (cx - w // 2 + skew, cy - h // 2),
        (cx + w // 2 + skew, cy - h // 2),
        (cx + w // 2 - skew, cy + h // 2),
        (cx - w // 2 - skew, cy + h // 2),
    ]
    draw.polygon(pts, fill=FG, outline=FG)
    draw_text(draw, cx, cy, text, 16, FONT_SM)


def rectangle(draw: ImageDraw.ImageDraw, cx: int, cy: int, text: str, w=160, h=62) -> None:
    x0, y0 = cx - w // 2, cy - h // 2
    draw.rectangle((x0, y0, x0 + w, y0 + h), fill=FG, outline=FG)
    draw_text(draw, cx, cy, text, 18, FONT_SM)


def predefined(draw: ImageDraw.ImageDraw, cx: int, cy: int, text: str, w=168, h=62) -> None:
    x0, y0 = cx - w // 2, cy - h // 2
    draw.rectangle((x0, y0, x0 + w, y0 + h), fill=FG, outline=FG)
    draw.line((x0 + 10, y0, x0 + 10, y0 + h), fill=INK, width=2)
    draw.line((x0 + w - 10, y0, x0 + w - 10, y0 + h), fill=INK, width=2)
    draw_text(draw, cx, cy, text, 16, FONT_SM)


def diamond(draw: ImageDraw.ImageDraw, cx: int, cy: int, text: str, w=150, h=92) -> None:
    pts = [(cx, cy - h // 2), (cx + w // 2, cy), (cx, cy + h // 2), (cx - w // 2, cy)]
    draw.polygon(pts, fill=FG, outline=FG)
    draw_text(draw, cx, cy, text, 14, FONT_SM)


def connector(draw: ImageDraw.ImageDraw, cx: int, cy: int, n: int) -> None:
    r = 22
    draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=FG, outline=FG)
    label = str(n)
    bbox = draw.textbbox((0, 0), label, font=FONT_CAPTION)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.text((cx - tw // 2, cy - th // 2 - 2), label, fill=INK, font=FONT_CAPTION)


def canvas(w: int, h: int) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (w, h), BG)
    return img, ImageDraw.Draw(img)


def save(img: Image.Image, name: str) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / name
    img.save(path, "PNG")
    return path


def fig_login() -> Path:
    img, d = canvas(1400, 720)
    oval(d, 110, 60, "Start")
    connector(d, 300, 60, 2)
    parallelogram(d, 170, 220, "Display Login Screen")
    parallelogram(d, 430, 220, "Enter username and password")
    diamond(d, 700, 220, "Is the Sign in button clicked")
    rectangle(d, 980, 220, "Validate the credentials")
    connector(d, 1220, 220, 1)
    connector(d, 80, 480, 1)
    diamond(d, 300, 480, "Are credentials valid")
    parallelogram(d, 300, 640, "Display error message")
    connector(d, 520, 640, 2)
    diamond(d, 620, 480, "Is the account an Administrator")
    predefined(d, 920, 420, "(Admin) Scanning Screen")
    predefined(d, 920, 540, "(User) Scanning Screen")
    oval(d, 1180, 480, "Stop")
    return save(img, "01-login.png")


def fig_scanning() -> Path:
    img, d = canvas(1500, 780)
    oval(d, 110, 50, "Start")
    parallelogram(d, 180, 180, "Display Scanning Screen")
    parallelogram(d, 460, 180, "Enter IP, hostname, CIDR, or range")
    diamond(d, 780, 180, "Is the Quick ping button clicked")
    diamond(d, 1080, 180, "Is the Deep nmap button clicked")
    connector(d, 1340, 180, 1)
    connector(d, 80, 420, 1)
    rectangle(d, 280, 420, "Run host discovery")
    parallelogram(d, 540, 420, "Display live hosts")
    parallelogram(d, 820, 420, "Select live hosts")
    diamond(d, 1100, 420, "Is Add to Inventory clicked")
    connector(d, 1360, 420, 2)
    connector(d, 80, 640, 2)
    rectangle(d, 280, 640, "Save selected hosts to SQLite assets")
    predefined(d, 560, 640, "Inventory Screen")
    oval(d, 820, 640, "Stop")
    return save(img, "02-scanning.png")


def fig_inventory() -> Path:
    img, d = canvas(1500, 860)
    oval(d, 110, 50, "Start")
    parallelogram(d, 200, 180, "Display Inventory Screen")
    parallelogram(d, 500, 180, "Filter, group by /24, and paginate")
    diamond(d, 820, 180, "Is the user an Administrator")
    parallelogram(d, 1120, 180, "View table only")
    connector(d, 1380, 180, 1)
    connector(d, 80, 400, 1)
    parallelogram(d, 260, 400, "Assign device and location")
    diamond(d, 540, 400, "Is Device types clicked")
    predefined(d, 820, 400, "Device types dialog")
    diamond(d, 1100, 400, "Is Locations clicked")
    predefined(d, 1360, 400, "Locations dialog")
    connector(d, 80, 640, 2)
    connector(d, 1380, 520, 2)
    diamond(d, 280, 640, "Is Check accessibility clicked")
    rectangle(d, 560, 640, "Probe WinRM then nmap MAC")
    diamond(d, 860, 640, "Is Delete selected clicked")
    rectangle(d, 1140, 640, "Delete selected assets")
    oval(d, 1380, 640, "Stop")
    return save(img, "03-inventory.png")


def fig_dashboard() -> Path:
    img, d = canvas(1500, 720)
    oval(d, 90, 80, "Start")
    parallelogram(d, 280, 80, "Display Dashboard Screen")
    parallelogram(d, 560, 80, "Show KPI cards")
    diamond(d, 860, 80, "Is View findings clicked")
    predefined(d, 1160, 80, "Findings Screen")
    connector(d, 1400, 80, 1)
    connector(d, 80, 300, 1)
    diamond(d, 280, 300, "Is a finding row clicked")
    predefined(d, 560, 300, "Asset Detail Screen")
    diamond(d, 860, 300, "Is an attention asset clicked")
    predefined(d, 1160, 300, "Asset Detail Screen")
    connector(d, 1400, 300, 2)
    connector(d, 80, 520, 2)
    diamond(d, 280, 520, "Is the User menu Log out clicked")
    rectangle(d, 560, 520, "Clear the in-memory session")
    predefined(d, 840, 520, "Login Screen")
    oval(d, 1100, 520, "Stop")
    return save(img, "04-dashboard.png")


def fig_asset_detail() -> Path:
    img, d = canvas(1400, 620)
    oval(d, 100, 70, "Start")
    parallelogram(d, 280, 180, "Display Asset Detail Screen")
    parallelogram(d, 560, 180, "View identity, ports, and ADHICS findings")
    diamond(d, 860, 180, "Is Open Scripts clicked")
    predefined(d, 1160, 180, "Scripts Screen")
    connector(d, 80, 420, 1)
    connector(d, 1320, 300, 1)
    diamond(d, 280, 420, "Is Back clicked")
    predefined(d, 560, 420, "Return to origin view")
    oval(d, 840, 420, "Stop")
    return save(img, "05-asset-detail.png")


def fig_findings() -> Path:
    img, d = canvas(1500, 700)
    oval(d, 100, 60, "Start")
    parallelogram(d, 260, 180, "Display Findings Screen")
    parallelogram(d, 540, 180, "Filter ADHICS gaps")
    diamond(d, 820, 180, "Is a finding row clicked")
    parallelogram(d, 1100, 180, "Display finding dialog")
    connector(d, 1360, 180, 1)
    connector(d, 80, 420, 1)
    diamond(d, 260, 420, "Is Acknowledge or Close clicked")
    rectangle(d, 540, 420, "Update finding status locally")
    diamond(d, 840, 420, "Is Open asset clicked")
    predefined(d, 1120, 420, "Asset Detail Screen")
    oval(d, 1360, 420, "Stop")
    return save(img, "06-findings.png")


def fig_adhics() -> Path:
    img, d = canvas(1500, 820)
    oval(d, 100, 50, "Start")
    parallelogram(d, 240, 170, "Display ADHICS catalog")
    diamond(d, 520, 170, "Is Add control clicked")
    parallelogram(d, 820, 170, "Enter ID, domain, description, status")
    rectangle(d, 1120, 170, "Save control in local catalog")
    connector(d, 1380, 170, 1)
    connector(d, 80, 400, 1)
    diamond(d, 260, 400, "Is a control row clicked")
    parallelogram(d, 540, 400, "Display control dialog")
    diamond(d, 820, 400, "Is Edit or Delete clicked")
    rectangle(d, 1100, 400, "Update or remove the control")
    connector(d, 1380, 400, 2)
    connector(d, 80, 640, 2)
    diamond(d, 280, 640, "Is Open Scripts clicked")
    predefined(d, 560, 640, "Scripts Screen")
    diamond(d, 860, 640, "Is a mapped asset clicked")
    predefined(d, 1140, 640, "Asset Detail Screen")
    oval(d, 1380, 640, "Stop")
    return save(img, "07-adhics.png")


def fig_scripts() -> Path:
    img, d = canvas(1500, 700)
    oval(d, 100, 60, "Start")
    parallelogram(d, 260, 180, "Display Scripts Screen")
    parallelogram(d, 540, 180, "Select a script")
    parallelogram(d, 820, 180, "Edit name, runner, timeout, and body")
    diamond(d, 1120, 180, "Is Save clicked")
    connector(d, 1380, 180, 1)
    connector(d, 80, 430, 1)
    rectangle(d, 260, 430, "Store script in local state")
    diamond(d, 540, 430, "Is Add script clicked")
    parallelogram(d, 820, 430, "Display Add script dialog")
    rectangle(d, 1100, 430, "Create script with Control ID and Run via")
    oval(d, 1360, 430, "Stop")
    return save(img, "08-scripts.png")


def fig_report() -> Path:
    img, d = canvas(1200, 420)
    oval(d, 100, 200, "Start")
    parallelogram(d, 320, 200, "Display Report preview")
    diamond(d, 580, 200, "Is Print clicked")
    rectangle(d, 840, 200, "Open the print dialog")
    oval(d, 1080, 200, "Stop")
    return save(img, "09-report.png")


def fig_settings() -> Path:
    img, d = canvas(1500, 760)
    oval(d, 100, 50, "Start")
    parallelogram(d, 260, 180, "Display Settings Screen")
    parallelogram(d, 540, 180, "Show company profile")
    diamond(d, 820, 180, "Is the user an Administrator")
    parallelogram(d, 1120, 180, "View profile only")
    connector(d, 1380, 180, 1)
    connector(d, 80, 400, 1)
    diamond(d, 260, 400, "Is Save profile clicked")
    rectangle(d, 540, 400, "Update company_profile in SQLite")
    diamond(d, 840, 400, "Is Add or Edit user clicked")
    parallelogram(d, 1140, 400, "Enter username, password, and level")
    connector(d, 1380, 400, 2)
    connector(d, 80, 620, 2)
    rectangle(d, 280, 620, "Save user as Administrator or User")
    diamond(d, 560, 620, "Is Delete user clicked")
    rectangle(d, 860, 620, "Delete account if not self or last admin")
    oval(d, 1140, 620, "Stop")
    return save(img, "10-settings.png")


FIGURES = [
    (
        fig_login,
        "Login Screen Flowchart",
        "This figure shows the Login Screen flowchart. The system displays the username and password fields. After the Sign in button is clicked, credentials are validated against SQLite. Invalid credentials return an error on the same screen. A successful login stores the session in the main process and opens the Scanning Screen for both Administrator and User levels.",
    ),
    (
        fig_scanning,
        "Scanning Screen Flowchart",
        "This figure shows the Scanning Screen flowchart. The user enters a single IP, hostname, CIDR, or IP range. Quick ping uses Windows ping -a. Deep nmap runs host discovery only. Live hosts remain in session memory until Add to Inventory is clicked. Selected hosts are written to SQLite; existing IPv4 rows are skipped. Both Administrator and User may scan and add hosts.",
    ),
    (
        fig_inventory,
        "Inventory Screen Flowchart",
        "This figure shows the Inventory Screen flowchart. Saved assets are listed with device and location filters, /24 grouping, and pagination. A User-level account may browse only. An Administrator may assign device and location, manage device types and locations in dialogs, run Check accessibility (WinRM then nmap MAC), and delete selected assets.",
    ),
    (
        fig_dashboard,
        "Dashboard Screen Flowchart",
        "This figure shows the Dashboard Screen flowchart. The prototype dashboard presents KPI cards, recent ADHICS findings, and assets needing attention using dummy data. View findings opens the Findings Screen. Clicking a finding or attention asset opens the dummy Asset Detail Screen. Log out from the user menu clears the in-memory session and returns to Login.",
    ),
    (
        fig_asset_detail,
        "Asset Detail Screen Flowchart",
        "This figure shows the Asset Detail Screen flowchart. This is a prototype drill-in from Dashboard, Findings, or ADHICS, not a live Inventory row. The screen shows dummy identity, ports, and ADHICS findings. Open Scripts jumps to the Scripts Screen for the mapped control. Back returns to the view that opened the detail.",
    ),
    (
        fig_findings,
        "Findings Screen Flowchart",
        "This figure shows the Findings Screen flowchart. Findings are ADHICS gaps rather than CVE records. The user may filter the list and open a finding dialog. Acknowledge and Close update local status only. Open asset navigates to the dummy Asset Detail Screen.",
    ),
    (
        fig_adhics,
        "ADHICS Screen Flowchart",
        "This figure shows the ADHICS Screen flowchart. The catalog is local application state with Control ID, domain, description, and status. Add, edit, and delete run in dialogs. Open Scripts uses the selected Control ID. A mapped dummy finding can open Asset Detail.",
    ),
    (
        fig_scripts,
        "Scripts Screen Flowchart",
        "This figure shows the Scripts Screen flowchart. Scripts are local stubs and are not executed. Each script has a Control ID from the ADHICS catalog and Run via WinRM or Nmap. The WinRM body uses Invoke-Command with {{ComputerName}}. Add script opens a dialog. Save stores the draft in renderer state.",
    ),
    (
        fig_report,
        "Report Screen Flowchart",
        "This figure shows the Report Screen flowchart. The prototype presents a print-like preview of clinic header, KPIs, ADHICS gaps, and an inventory excerpt using dummy data. Print opens the system print dialog. No PDF file is written.",
    ),
    (
        fig_settings,
        "Settings Screen Flowchart",
        "This figure shows the Settings Screen flowchart. Both levels can view the company profile. Only an Administrator can save name, address, contact, and notes to SQLite, and can add, edit, or delete local users with Administrator or User level. The signed-in account cannot be deleted, and the last administrator cannot be removed.",
    ),
]


def set_run_font(run, size=12, bold=False, italic=False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_justified_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    set_run_font(run, 12)


def build_doc(paths: list[tuple[Path, str, str]]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = heading.add_run("Flowchart")
    set_run_font(run, 12, bold=True)

    for index, (image, caption, body) in enumerate(paths, start=1):
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(image), width=Inches(6.4))

        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(f"Figure {index}. {caption}")
        set_run_font(run, 12, bold=True)

        add_justified_paragraph(doc, body)
        if index < len(paths):
            doc.add_paragraph()

    doc.save(str(DOCX_PATH))


def main() -> None:
    built: list[tuple[Path, str, str]] = []
    for fn, caption, body in FIGURES:
        built.append((fn(), caption, body))
    build_doc(built)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
